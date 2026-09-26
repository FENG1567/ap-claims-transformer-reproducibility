from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd


ROOT = Path(__file__).parents[2]
REFERENCE = ROOT / "work" / "reference"
OUT = ROOT / "outputs" / "stage2_lock" / "planned_readmission"
CATALOG = REFERENCE / "official_icd10_catalog_fy2018_2023.parquet"


@dataclass(frozen=True)
class AnnualSource:
    calendar_year: int
    reporting_year: int
    fiscal_years: tuple[int, int]
    specification: Path
    mapping: Path


SOURCES = (
    AnnualSource(
        2018,
        2020,
        (2018, 2019),
        REFERENCE / "CMS_annual_HWR" / "2020_HWR_SuppFile.xlsx",
        REFERENCE
        / "CMS_2021_HWR_supplement"
        / "YaleModified_CCS_PCS_CM_Map_v2020_final508.xlsx",
    ),
    AnnualSource(
        2019,
        2021,
        (2019, 2020),
        REFERENCE / "CMS_2021_HWR_supplement" / "2021_HWR.xlsx",
        REFERENCE
        / "CMS_2021_HWR_supplement"
        / "YaleModified_CCS_PCS_CM_Map_v2020_final508.xlsx",
    ),
    AnnualSource(
        2020,
        2022,
        (2020, 2021),
        REFERENCE / "CMS_annual_HWR" / "2022_HWR (3).xlsx",
        REFERENCE
        / "CMS_annual_HWR"
        / "YaleMod_CCS_PCS_CM_Map_for_2022-external_v1.0.xlsx",
    ),
    AnnualSource(
        2021,
        2023,
        (2021, 2022),
        REFERENCE / "CMS_annual_HWR" / "2023_HWR_v1.0.xlsx",
        REFERENCE
        / "CMS_annual_HWR"
        / "YaleMod_CCS_CM_PCS_Map_for_2023_Reporting_External_9.29.22.xlsx",
    ),
)


TABLES = {
    "PR.1": {"kind": "category", "system": "PCS", "position": "any_procedure"},
    "PR.2": {"kind": "category", "system": "CM", "position": "principal_diagnosis"},
    "PR.3": {"kind": "mixed", "system": "PCS", "position": "any_procedure"},
    "PR.4": {"kind": "mixed", "system": "CM", "position": "principal_diagnosis"},
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_code(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().upper().replace(".", "")


def normalize_category(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def normalize_status(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def sheet_for_prefix(workbook: openpyxl.Workbook, prefix: str):
    matches = [name for name in workbook.sheetnames if name.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one {prefix} sheet, found {matches}")
    return workbook[matches[0]]


def read_specification(source: AnnualSource) -> pd.DataFrame:
    workbook = openpyxl.load_workbook(
        source.specification, read_only=True, data_only=True
    )
    records: list[dict[str, object]] = []
    for table, metadata in TABLES.items():
        sheet = sheet_for_prefix(workbook, table)
        blank_run = 0
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=3, min_col=1, max_col=6, values_only=True), start=3
        ):
            values = tuple(row[:6])
            if not any(value is not None for value in values):
                blank_run += 1
                if blank_run >= 500:
                    break
                continue
            blank_run = 0
            first = str(values[0]).strip() if values[0] is not None else ""
            if first.lower().startswith("end of worksheet"):
                break
            if metadata["kind"] == "category":
                record_type = "CCS"
                value = normalize_category(values[0])
                description = str(values[1] or "").strip()
                associated_ccs = value
                status = normalize_status(values[2])
                notes = str(values[3] or "").strip()
            else:
                record_type = first.upper()
                value = (
                    normalize_category(values[1])
                    if record_type == "CCS"
                    else normalize_code(values[1])
                )
                description = str(values[2] or "").strip()
                associated_ccs = str(values[3] or "").strip()
                status = normalize_status(values[4])
                notes = str(values[5] or "").strip()
            if not value:
                raise RuntimeError(
                    f"Blank value in {source.specification.name} {sheet.title} row {row_number}"
                )
            if record_type not in {"CCS", "ICD-10-CM", "ICD-10-PCS"}:
                raise RuntimeError(
                    f"Unexpected record type {record_type!r} in "
                    f"{source.specification.name} {sheet.title} row {row_number}"
                )
            records.append(
                {
                    "calendar_year": source.calendar_year,
                    "reporting_year": source.reporting_year,
                    "table": table,
                    "record_type": record_type,
                    "value": value,
                    "description": description,
                    "associated_ccs": associated_ccs,
                    "status": status,
                    "active": not status.lower().startswith("remove"),
                    "notes": notes,
                    "source_sheet": sheet.title,
                    "source_row": row_number,
                }
            )
    workbook.close()
    result = pd.DataFrame(records)
    duplicates = result.duplicated(
        ["calendar_year", "table", "record_type", "value"], keep=False
    )
    if duplicates.any():
        sample = result.loc[
            duplicates, ["calendar_year", "table", "record_type", "value"]
        ].head(20)
        raise RuntimeError(f"Duplicate specification rows:\n{sample}")
    return result


def read_mapping(source: AnnualSource) -> pd.DataFrame:
    workbook = openpyxl.load_workbook(source.mapping, read_only=True, data_only=True)
    frames: list[pd.DataFrame] = []
    for system, sheet_name in (("CM", "CCS-ICD10CM"), ("PCS", "CCS-ICD10PCS")):
        sheet = workbook[sheet_name]
        records: list[dict[str, object]] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=3, min_col=1, max_col=5, values_only=True), start=3
        ):
            values = tuple(row[:5])
            if not any(value is not None for value in values):
                continue
            if str(values[0] or "").strip().lower().startswith("end of worksheet"):
                break
            code = normalize_code(values[0])
            category = normalize_category(values[2])
            if not code or not category:
                raise RuntimeError(
                    f"Incomplete mapping in {source.mapping.name} {sheet_name} row {row_number}"
                )
            records.append(
                {
                    "calendar_year": source.calendar_year,
                    "code_system": system,
                    "code": code,
                    "description": str(values[1] or "").strip(),
                    "ccs_category": category,
                    "ccs_description": str(values[3] or "").strip(),
                    "mapping_notes": str(values[4] or "").strip(),
                }
            )
        frame = pd.DataFrame(records)
        duplicated = frame.duplicated(["code_system", "code"], keep=False)
        if duplicated.any():
            inconsistent = (
                frame.loc[duplicated]
                .groupby(["code_system", "code"])["ccs_category"]
                .nunique()
            )
            inconsistent = inconsistent[inconsistent > 1]
            if not inconsistent.empty:
                raise RuntimeError(
                    f"Codes map to multiple CCS categories in {source.mapping.name}: "
                    f"{inconsistent.head(20).to_dict()}"
                )
            frame = frame.drop_duplicates(["code_system", "code"], keep="first")
        frames.append(frame)
    workbook.close()
    return pd.concat(frames, ignore_index=True)


def build_valid_catalog() -> pd.DataFrame:
    catalog = pd.read_parquet(CATALOG)
    catalog = catalog[catalog["billable"].fillna(False)].copy()
    catalog["code_system"] = catalog["system"].astype(str)
    catalog["code"] = catalog["code"].map(normalize_code)
    return catalog[["fiscal_year", "code_system", "code"]].drop_duplicates()


def expand_code_sets(
    specifications: pd.DataFrame,
    mappings: pd.DataFrame,
    valid_catalog: pd.DataFrame,
) -> tuple[pd.DataFrame, dict[str, object]]:
    output: list[pd.DataFrame] = []
    qc: dict[str, object] = {"years": {}, "status_rule": "active unless status starts with remove"}
    for source in SOURCES:
        year = source.calendar_year
        spec = specifications[
            (specifications.calendar_year == year) & specifications.active
        ].copy()
        mapping = mappings[mappings.calendar_year == year].copy()
        valid = valid_catalog[
            valid_catalog.fiscal_year.isin(source.fiscal_years)
        ][["code_system", "code"]].drop_duplicates()
        valid_sets = {
            system: set(group.code)
            for system, group in valid.groupby("code_system", sort=False)
        }
        mapping_sets = {
            system: set(group.code)
            for system, group in mapping.groupby("code_system", sort=False)
        }
        year_qc: dict[str, object] = {
            "reporting_year": source.reporting_year,
            "fiscal_years": list(source.fiscal_years),
            "valid_catalog_codes": {
                system: len(codes) for system, codes in sorted(valid_sets.items())
            },
            "mapping_codes": {
                system: len(codes) for system, codes in sorted(mapping_sets.items())
            },
            "valid_codes_missing_from_mapping": {
                system: len(valid_sets.get(system, set()) - mapping_sets.get(system, set()))
                for system in ("CM", "PCS")
            },
            "tables": {},
        }
        for table, metadata in TABLES.items():
            system = str(metadata["system"])
            table_spec = spec[spec.table == table]
            categories = set(
                table_spec.loc[table_spec.record_type == "CCS", "value"].astype(str)
            )
            direct_type = "ICD-10-CM" if system == "CM" else "ICD-10-PCS"
            direct = set(
                table_spec.loc[table_spec.record_type == direct_type, "value"].astype(str)
            )
            expanded = set(
                mapping.loc[
                    (mapping.code_system == system)
                    & mapping.ccs_category.isin(categories),
                    "code",
                ]
            )
            combined = expanded | direct
            valid_combined = combined & valid_sets.get(system, set())
            invalid_direct = direct - valid_sets.get(system, set())
            rows = pd.DataFrame(
                {
                    "calendar_year": year,
                    "reporting_year": source.reporting_year,
                    "table": table,
                    "code_system": system,
                    "code": sorted(valid_combined),
                }
            )
            output.append(rows)
            year_qc["tables"][table] = {
                "active_ccs_categories": len(categories),
                "active_direct_codes": len(direct),
                "expanded_category_codes": len(expanded),
                "final_valid_codes": len(valid_combined),
                "direct_codes_invalid_for_calendar_year": len(invalid_direct),
                "direct_invalid_examples": sorted(invalid_direct)[:20],
            }
        qc["years"][str(year)] = year_qc
    code_sets = pd.concat(output, ignore_index=True)
    if code_sets.duplicated(["calendar_year", "table", "code_system", "code"]).any():
        raise RuntimeError("Expanded code sets contain duplicate rows")
    return code_sets, qc


def write_algorithm_rules() -> None:
    (OUT / "algorithm_rules.yml").write_text(
        """version: PRA_v4.0_annual_alignment_v1
status: LOCKED_PRE_2022
logic:
  planned_if:
    - any_procedure_in_PR1
    - principal_diagnosis_in_PR2
    - any_procedure_in_PR3 AND principal_diagnosis_not_in_PR4
  unplanned_if: NOT planned
positions:
  PR1: any_procedure_position
  PR2: principal_diagnosis_only
  PR3: any_procedure_position
  PR4: principal_diagnosis_only
annual_alignment:
  2018: CMS_HWR_2020_reporting_specification_FY2018_FY2019_codes
  2019: CMS_HWR_2021_reporting_specification_FY2019_FY2020_codes
  2020: CMS_HWR_2022_reporting_specification_FY2020_FY2021_codes
  2021: CMS_HWR_2023_reporting_specification_FY2021_FY2022_codes
status_interpretation:
  active: blank_dash_added_revised_or_other_non_remove_status
  inactive: status_normalized_starts_with_remove
unknown_mapping_policy:
  primary: flag_algorithm_unknown_and_exclude_if_an_observed_code_is_absent_from_the_year_map_and_not_an_explicit_direct_code
  sensitivity:
    - treat_all_unknown_as_unplanned
    - treat_all_unknown_as_planned
2022_policy: sealed_until_stage7_then_apply_predeclared_annual_alignment_rule
""",
        encoding="utf-8",
    )


def main() -> None:
    required = {CATALOG}
    for source in SOURCES:
        required.update({source.specification, source.mapping})
    missing = sorted(str(path) for path in required if not path.is_file())
    if missing:
        raise RuntimeError(f"Missing planned-readmission references: {missing}")
    OUT.mkdir(parents=True, exist_ok=True)

    specifications = pd.concat(
        [read_specification(source) for source in SOURCES], ignore_index=True
    )
    mappings = pd.concat([read_mapping(source) for source in SOURCES], ignore_index=True)
    catalog = build_valid_catalog()
    code_sets, qc = expand_code_sets(specifications, mappings, catalog)

    specifications.to_csv(OUT / "annual_spec_rows.csv", index=False, lineterminator="\n")
    mappings.to_parquet(
        OUT / "annual_ccs_mapping.parquet", index=False, compression="zstd"
    )
    code_sets.to_parquet(
        OUT / "annual_algorithm_code_sets.parquet", index=False, compression="zstd"
    )
    write_algorithm_rules()

    source_manifest = {
        str(path.relative_to(ROOT)): {"bytes": path.stat().st_size, "sha256": sha256(path)}
        for path in sorted(required)
    }
    artifacts = [
        OUT / "annual_spec_rows.csv",
        OUT / "annual_ccs_mapping.parquet",
        OUT / "annual_algorithm_code_sets.parquet",
        OUT / "algorithm_rules.yml",
    ]
    lock = {
        "status": "PASS",
        "algorithm": "CMS Planned Readmission Algorithm Version 4.0",
        "years_locked": [2018, 2019, 2020, 2021],
        "sealed_year": 2022,
        "qc": qc,
        "sources": source_manifest,
        "artifacts": {
            str(path.relative_to(OUT)): {
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
            for path in artifacts
        },
    }
    (OUT / "planned_readmission_lock.json").write_text(
        json.dumps(lock, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "status": "PASS",
                "specification_rows": len(specifications),
                "mapping_rows": len(mappings),
                "algorithm_code_rows": len(code_sets),
                "lock_sha256": sha256(OUT / "planned_readmission_lock.json"),
                "qc": qc,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
